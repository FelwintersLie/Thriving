from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

GRID_START_MINUTE = 7 * 60 + 30
GRID_END_MINUTE = 18 * 60
GRID_SLOT_MINUTES = 15


DEFAULT_STYLE: Dict[str, Any] = {
    "time_col_w": 85,
    "header_h": 42,
    "row_h": 22,
    "col_w": 180,
    "header_time_fill": "#0b4f6c",
    "header_col_fill": "#1d3557",
    "header_col_outline": "#f1faee",
    "time_even_fill": "#f8f9fa",
    "time_odd_fill": "#e9ecef",
    "time_outline": "#adb5bd",
    "cell_fill": "#ffffff",
    "cell_outline": "#dee2e6",
}


def grid_minutes(start_minute: int = GRID_START_MINUTE, end_minute: int = GRID_END_MINUTE, slot_minutes: int = GRID_SLOT_MINUTES) -> List[int]:
    return list(range(start_minute, end_minute, slot_minutes))


def _to_ampm(minute: int) -> str:
    hour_24 = minute // 60
    minute_of_hour = minute % 60
    period = "AM" if hour_24 < 12 else "PM"
    hour_12 = hour_24 % 12 or 12
    return f"{hour_12}:{minute_of_hour:02d} {period}"


def build_schedule_layout_model(
    appointments: List[Dict[str, Any]],
    planning_dates: List[str],
    discipline_colors: Dict[str, str],
    *,
    grid_mode: str = "Patient Grid",
    program_filter: str = "Both",
    selected_request_id: Optional[str] = None,
    visible_dates: Optional[List[str]] = None,
) -> Dict[str, Any]:
    minutes = grid_minutes()
    style = dict(DEFAULT_STYLE)
    filtered = list(appointments)

    if visible_dates:
        allowed = set(visible_dates)
        filtered = [a for a in filtered if a.get("date_key") in allowed]

    if program_filter in {"IOP", "EVAL"}:
        filtered = [a for a in filtered if a.get("program_type", "IOP") == program_filter]

    date_keys_in_use = {a.get("date_key") for a in filtered if a.get("date_key")}
    ordered_dates = [d for d in planning_dates if d in date_keys_in_use] + sorted(date_keys_in_use.difference(planning_dates))
    if not ordered_dates and visible_dates:
        ordered_dates = list(visible_dates)

    layout: Dict[str, Any] = {
        "canvas": {"width": 500, "height": 300},
        "rectangles": [],
        "texts": [],
        "legend_disciplines": [],
        "metadata": {
            "grid_mode": grid_mode,
            "program_filter": program_filter,
            "ordered_dates": ordered_dates,
            "appointment_count": len(filtered),
        },
        "style": style,
    }

    if not filtered:
        layout["texts"].append(
            {
                "x": 16,
                "y": 20,
                "text": "No appointments scheduled for this view/filter.",
                "fill": "#003049",
                "font_size": 11,
                "bold": True,
                "anchor": "w",
                "justify": "left",
            }
        )
        return layout

    if grid_mode == "Room Grid":
        axis_labels = sorted({a.get("room", "") for a in filtered if a.get("room")})
        label_prefix = "Room"
    elif grid_mode == "Provider Grid":
        axis_labels = sorted({a.get("provider", "") for a in filtered if a.get("provider")})
        label_prefix = "Provider"
    else:
        axis_labels = sorted(
            {pid for a in filtered for pid in a.get("patients", [])},
            key=lambda x: (x[:1], int(x[1:]) if x[1:].isdigit() else x),
        )
        label_prefix = "Patient"

    if not axis_labels:
        layout["texts"].append(
            {
                "x": 16,
                "y": 20,
                "text": "No axis labels available for current mode.",
                "fill": "#003049",
                "font_size": 11,
                "bold": True,
                "anchor": "w",
                "justify": "left",
            }
        )
        return layout

    time_col_w = style["time_col_w"]
    header_h = style["header_h"]
    row_h = style["row_h"]
    col_w = style["col_w"]

    col_pairs = [(d, p) for d in ordered_dates for p in axis_labels]
    total_w = time_col_w + len(col_pairs) * col_w
    total_h = header_h + len(minutes) * row_h
    layout["canvas"] = {"width": total_w, "height": total_h}

    layout["rectangles"].append({"x0": 0, "y0": 0, "x1": time_col_w, "y1": header_h, "fill": style["header_time_fill"], "outline": style["header_time_fill"], "width": 1})
    layout["texts"].append({"x": time_col_w / 2, "y": header_h / 2, "x0": 0, "y0": 0, "x1": time_col_w, "y1": header_h, "text": "Time", "fill": "#ffffff", "font_size": 10, "bold": True, "anchor": "center", "justify": "center"})

    for idx, (date_key, axis_val) in enumerate(col_pairs):
        x0 = time_col_w + idx * col_w
        x1 = x0 + col_w
        layout["rectangles"].append({"x0": x0, "y0": 0, "x1": x1, "y1": header_h, "fill": style["header_col_fill"], "outline": style["header_col_outline"], "width": 1})
        layout["texts"].append({"x": (x0 + x1) / 2, "y": header_h / 2, "x0": x0 + 4, "y0": 2, "x1": x1 - 4, "y1": header_h - 2, "text": f"{date_key}\n{label_prefix} {axis_val}", "fill": "#ffffff", "font_size": 8, "bold": True, "anchor": "center", "justify": "center"})

    for row_idx, minute in enumerate(minutes):
        y0 = header_h + row_idx * row_h
        y1 = y0 + row_h
        time_fill = style["time_even_fill"] if row_idx % 2 == 0 else style["time_odd_fill"]
        layout["rectangles"].append({"x0": 0, "y0": y0, "x1": time_col_w, "y1": y1, "fill": time_fill, "outline": style["time_outline"], "width": 1})
        layout["texts"].append({"x": time_col_w / 2, "y": (y0 + y1) / 2, "x0": 2, "y0": y0, "x1": time_col_w - 2, "y1": y1, "text": _to_ampm(minute), "fill": "#1b263b", "font_size": 8, "bold": False, "anchor": "center", "justify": "center"})
        for col_idx in range(len(col_pairs)):
            x0 = time_col_w + col_idx * col_w
            x1 = x0 + col_w
            layout["rectangles"].append({"x0": x0, "y0": y0, "x1": x1, "y1": y1, "fill": style["cell_fill"], "outline": style["cell_outline"], "width": 1})

    pair_index = {pair: idx for idx, pair in enumerate(col_pairs)}
    used_disciplines = set()
    for appt in filtered:
        start_idx = max(0, (int(appt["start"]) - GRID_START_MINUTE) // GRID_SLOT_MINUTES)
        end_idx = min(len(minutes), (int(appt["end"]) - GRID_START_MINUTE) // GRID_SLOT_MINUTES)
        if end_idx <= start_idx:
            continue
        if grid_mode == "Room Grid":
            keys = [appt.get("room", "")]
        elif grid_mode == "Provider Grid":
            keys = [appt.get("provider", "")]
        else:
            keys = list(appt.get("patients", []))

        for axis_key in keys:
            pair = (appt["date_key"], axis_key)
            if pair not in pair_index:
                continue
            col_idx = pair_index[pair]
            x0 = time_col_w + col_idx * col_w + 1
            x1 = x0 + col_w - 2
            y0 = header_h + start_idx * row_h + 1
            y1 = header_h + end_idx * row_h - 1
            is_selected = bool(selected_request_id and appt.get("request_id") == selected_request_id)
            layout["rectangles"].append(
                {
                    "x0": x0,
                    "y0": y0,
                    "x1": x1,
                    "y1": y1,
                    "fill": discipline_colors.get(appt.get("discipline", ""), "#ffb3c1"),
                    "outline": "#d00000" if is_selected else "#495057",
                    "width": 3 if is_selected else 2,
                    "kind": "appointment",
                    "request_id": appt.get("request_id", ""),
                }
            )
            layout["texts"].append(
                {
                    "x": (x0 + x1) / 2,
                    "y": (y0 + y1) / 2,
                    "x0": x0 + 3,
                    "y0": y0 + 2,
                    "x1": x1 - 3,
                    "y1": y1 - 2,
                    "text": f"{appt.get('discipline', '')}\n{appt.get('room', '')}\n{appt.get('provider', '')}\n{appt.get('program_type', 'IOP')}",
                    "fill": "#1b263b",
                    "font_size": 8,
                    "bold": False,
                    "anchor": "center",
                    "justify": "center",
                    "kind": "appointment",
                    "request_id": appt.get("request_id", ""),
                }
            )
            used_disciplines.add(appt.get("discipline", ""))

    layout["legend_disciplines"] = sorted(d for d in used_disciplines if d)
    return layout


def draw_layout_on_tk_canvas(canvas: Any, layout: Dict[str, Any]) -> None:
    canvas.delete("all")
    for rect in layout.get("rectangles", []):
        tags = ()
        if rect.get("kind") == "appointment":
            tags = ("appointment", f"req:{rect.get('request_id', '')}")
        canvas.create_rectangle(
            rect["x0"],
            rect["y0"],
            rect["x1"],
            rect["y1"],
            fill=rect.get("fill", ""),
            outline=rect.get("outline", "#000000"),
            width=rect.get("width", 1),
            tags=tags,
        )
    for text in layout.get("texts", []):
        tags = ()
        if text.get("kind") == "appointment":
            tags = ("appointment", f"req:{text.get('request_id', '')}")
        canvas.create_text(
            text["x"],
            text["y"],
            text=text.get("text", ""),
            fill=text.get("fill", "#000000"),
            font=("Segoe UI", text.get("font_size", 8), "bold" if text.get("bold") else "normal"),
            justify=text.get("justify", "center"),
            anchor=text.get("anchor", "center"),
            tags=tags,
        )
    c = layout.get("canvas", {})
    canvas.config(scrollregion=(0, 0, int(c.get("width", 500)), int(c.get("height", 300))))


def render_layout_to_png(layout: Dict[str, Any], out_path: Path) -> None:
    try:
        from PIL import Image, ImageDraw
    except ImportError as exc:
        raise RuntimeError("Pillow is required for PNG export. Install with: python3 -m pip install Pillow") from exc

    canvas = layout.get("canvas", {})
    width = max(1, int(canvas.get("width", 500)))
    height = max(1, int(canvas.get("height", 300)))
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    for rect in layout.get("rectangles", []):
        draw.rectangle([(rect["x0"], rect["y0"]), (rect["x1"], rect["y1"])], fill=rect.get("fill", "white"), outline=rect.get("outline", "#000000"), width=int(rect.get("width", 1)))

    for text in layout.get("texts", []):
        content = str(text.get("text", ""))
        x = float(text.get("x", 0))
        y = float(text.get("y", 0))
        anchor = text.get("anchor", "center")
        if anchor == "w":
            draw.multiline_text((x, y), content, fill=text.get("fill", "#000000"), anchor="la", align="left", spacing=2)
        else:
            draw.multiline_text((x, y), content, fill=text.get("fill", "#000000"), anchor="mm", align=text.get("justify", "center"), spacing=2)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(out_path, format="PNG")


def _hex_to_rgb(hex_color: str) -> Tuple[int, int, int]:
    clean = hex_color.strip().lstrip("#")
    if len(clean) != 6:
        return 0, 0, 0
    return int(clean[0:2], 16), int(clean[2:4], 16), int(clean[4:6], 16)


def _truncate_with_ellipsis(text: str, max_chars: int) -> str:
    clean = (text or "").strip()
    if len(clean) <= max_chars:
        return clean
    if max_chars <= 1:
        return "…"
    return clean[: max_chars - 1].rstrip() + "…"


def _fit_text_for_box(text: str, width_px: float, height_px: float, font_px: float) -> str:
    approx_char_w = max(4.0, font_px * 0.62)
    max_chars_per_line = max(4, int(width_px / approx_char_w))
    max_lines = max(1, int(height_px / max(10.0, font_px * 1.25)))
    lines: List[str] = []
    for raw_line in str(text or "").splitlines()[:max_lines]:
        lines.append(_truncate_with_ellipsis(raw_line, max_chars_per_line))
    while len(lines) < max_lines:
        break
    if len(str(text or "").splitlines()) > max_lines and lines:
        lines[-1] = _truncate_with_ellipsis(lines[-1], max_chars_per_line)
    return "\n".join(lines)



def _hex_to_argb(hex_color: str) -> str:
    clean = hex_color.strip().lstrip("#")
    if len(clean) != 6:
        clean = "000000"
    return "FF" + clean.upper()


def export_layout_to_xlsx(layout: Dict[str, Any], out_path: Path, title: str = "Schedule") -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel export. Install with: python3 -m pip install openpyxl") from exc

    style = layout.get("style", {})
    canvas = layout.get("canvas", {})
    time_col_w = int(style.get("time_col_w", 85))
    header_h = int(style.get("header_h", 42))
    row_h = int(style.get("row_h", 22))
    col_w = int(style.get("col_w", 180))
    total_w = int(canvas.get("width", 500))
    total_h = int(canvas.get("height", 300))

    time_rows = max(1, int(round((total_h - header_h) / max(1, row_h))))
    data_cols = max(1, int(round((total_w - time_col_w) / max(1, col_w))))

    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule"

    thin = Side(style="thin", color="FFADB5BD")
    base_border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.column_dimensions["A"].width = max(8, time_col_w / 7.0)
    for c in range(2, data_cols + 2):
        ws.column_dimensions[get_column_letter(c)].width = max(10, col_w / 7.0)

    ws.row_dimensions[1].height = max(18, header_h * 0.75)
    for r in range(2, time_rows + 2):
        ws.row_dimensions[r].height = max(14, row_h * 0.75)

    for r in range(1, time_rows + 2):
        for c in range(1, data_cols + 2):
            cell = ws.cell(row=r, column=c)
            cell.border = base_border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for rect in layout.get("rectangles", []):
        x0 = int(rect.get("x0", 0))
        y0 = int(rect.get("y0", 0))
        x1 = int(rect.get("x1", 0))
        y1 = int(rect.get("y1", 0))
        if y1 <= 0:
            row = 1
            if x1 <= time_col_w:
                c0 = c1 = 1
            else:
                c0 = max(2, 2 + int((x0 - time_col_w) // max(1, col_w)))
                c1 = max(c0, 1 + int((x1 - time_col_w) // max(1, col_w)))
            for c in range(c0, min(c1, data_cols + 1) + 1):
                ws.cell(row=row, column=c).fill = PatternFill(fill_type="solid", fgColor=_hex_to_argb(rect.get("fill", "#ffffff")))
            continue

        if x1 <= time_col_w:
            c0 = c1 = 1
        else:
            c0 = max(2, 2 + int((x0 - time_col_w) // max(1, col_w)))
            c1 = max(c0, 1 + int((x1 - time_col_w) // max(1, col_w)))

        r0 = 2 + max(0, int((y0 - header_h) // max(1, row_h)))
        r1 = 1 + max(1, int((y1 - header_h) // max(1, row_h)))

        for rr in range(r0, min(r1, time_rows + 1) + 1):
            for cc in range(c0, min(c1, data_cols + 1) + 1):
                ws.cell(row=rr, column=cc).fill = PatternFill(fill_type="solid", fgColor=_hex_to_argb(rect.get("fill", "#ffffff")))

    appointment_text_by_id: Dict[str, str] = {}
    for text in layout.get("texts", []):
        if text.get("kind") == "appointment" and text.get("request_id"):
            appointment_text_by_id[str(text.get("request_id"))] = str(text.get("text", ""))

    for text in layout.get("texts", []):
        content = str(text.get("text", "")).strip()
        if not content:
            continue
        x = float(text.get("x", 0))
        y = float(text.get("y", 0))
        if y <= header_h:
            row = 1
            col = 1 if x <= time_col_w else 2 + int((x - time_col_w) // max(1, col_w))
            cell = ws.cell(row=row, column=max(1, min(data_cols + 1, col)))
            cell.value = content
            cell.font = Font(bold=bool(text.get("bold", False)), color=_hex_to_argb(text.get("fill", "#000000")))
            continue
        if x <= time_col_w:
            row = 2 + int((y - header_h) // max(1, row_h))
            if 2 <= row <= time_rows + 1:
                cell = ws.cell(row=row, column=1)
                cell.value = content
                cell.font = Font(bold=bool(text.get("bold", False)), color=_hex_to_argb(text.get("fill", "#000000")))

    for rect in layout.get("rectangles", []):
        if rect.get("kind") != "appointment":
            continue
        req_id = str(rect.get("request_id", ""))
        content = appointment_text_by_id.get(req_id, "")

        x0 = int(rect.get("x0", 0))
        x1 = int(rect.get("x1", 0))
        y0 = int(rect.get("y0", 0))
        y1 = int(rect.get("y1", 0))

        col = max(2, 2 + int((x0 - time_col_w) // max(1, col_w)))
        row_start = 2 + max(0, int((y0 - header_h) // max(1, row_h)))
        row_end = 1 + max(1, int((y1 - header_h) // max(1, row_h)))
        row_start = max(2, min(time_rows + 1, row_start))
        row_end = max(row_start, min(time_rows + 1, row_end))

        if row_end > row_start:
            ws.merge_cells(start_row=row_start, start_column=col, end_row=row_end, end_column=col)
        cell = ws.cell(row=row_start, column=col)
        cell.value = content
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.font = Font(color="FF1B263B")
        cell.fill = PatternFill(fill_type="solid", fgColor=_hex_to_argb(rect.get("fill", "#ffffff")))

    if title:
        ws["A1"].comment = None
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))

def export_layout_to_pptx(layout: Dict[str, Any], out_path: Path, title: str) -> None:
    try:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt
    except ImportError as exc:
        raise RuntimeError("python-pptx is required for PowerPoint export. Install with: python3 -m pip install python-pptx") from exc

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    title_box = slide.shapes.add_textbox(Inches(0.3), Inches(0.1), Inches(12.7), Inches(0.4))
    title_tf = title_box.text_frame
    title_tf.text = title
    title_tf.paragraphs[0].font.size = Pt(16)
    title_tf.paragraphs[0].font.bold = True

    canvas = layout.get("canvas", {})
    width = float(canvas.get("width", 500))
    height = float(canvas.get("height", 300))
    left = Inches(0.3)
    top = Inches(0.7)
    max_w = Inches(12.7)
    max_h = Inches(6.6)
    scale = min(float(max_w) / max(width, 1.0), float(max_h) / max(height, 1.0))

    for rect in layout.get("rectangles", []):
        x = left + int(rect["x0"] * scale)
        y = top + int(rect["y0"] * scale)
        w = max(1, int((rect["x1"] - rect["x0"]) * scale))
        h = max(1, int((rect["y1"] - rect["y0"]) * scale))
        shp = slide.shapes.add_shape(1, x, y, w, h)
        fr, fg, fb = _hex_to_rgb(rect.get("fill", "#ffffff"))
        or_, og, ob = _hex_to_rgb(rect.get("outline", "#000000"))
        shp.fill.solid()
        shp.fill.fore_color.rgb = RGBColor(fr, fg, fb)
        shp.line.color.rgb = RGBColor(or_, og, ob)
        shp.line.width = Pt(max(0.75, float(rect.get("width", 1))))

    for text in layout.get("texts", []):
        content = str(text.get("text", "")).strip()
        if not content:
            continue
        x = left + int((float(text.get("x", 0)) - 50) * scale)
        y = top + int((float(text.get("y", 0)) - 14) * scale)
        box_w = int(max(30, (text.get("x1", text.get("x", 0) + 50) - text.get("x0", text.get("x", 0) - 50)) * scale))
        box_h = int(max(16, (text.get("y1", text.get("y", 0) + 14) - text.get("y0", text.get("y", 0) - 14)) * scale))
        tb = slide.shapes.add_textbox(x, y, box_w, box_h)
        tf = tb.text_frame
        tf.clear()
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = _fit_text_for_box(content, box_w, box_h, max(7, int(float(text.get("font_size", 8)) * scale * 1.8)))
        p.font.size = Pt(max(7, int(float(text.get("font_size", 8)) * scale * 1.8)))
        p.font.bold = bool(text.get("bold", False))
        p.alignment = PP_ALIGN.LEFT if text.get("anchor") == "w" else PP_ALIGN.CENTER

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
