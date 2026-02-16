import importlib.util
import tempfile
import unittest
from pathlib import Path

from app.schedule_exports import build_schedule_layout_model, export_layout_to_pptx, render_layout_to_png
from app.windows_gui import DISCIPLINE_COLORS


class ScheduleExportTests(unittest.TestCase):
    def _sample_layout(self):
        appointments = [
            {
                "request_id": "req_1",
                "date_key": "2026-01-05",
                "start": 510,
                "end": 570,
                "discipline": "Physical Therapy",
                "provider": "Provider A",
                "room": "Room 1",
                "patients": ["I1"],
                "program_type": "IOP",
            },
            {
                "request_id": "req_2",
                "date_key": "2026-01-05",
                "start": 600,
                "end": 660,
                "discipline": "Neuropsychology",
                "provider": "Provider B",
                "room": "Room 2",
                "patients": ["E1"],
                "program_type": "EVAL",
            },
        ]
        return build_schedule_layout_model(
            appointments,
            planning_dates=["2026-01-05", "2026-01-06"],
            discipline_colors=DISCIPLINE_COLORS,
            grid_mode="Patient Grid",
            program_filter="Both",
        )

    def test_layout_model_is_deterministic(self):
        a = self._sample_layout()
        b = self._sample_layout()
        self.assertEqual(a, b)
        self.assertGreater(len(a.get("rectangles", [])), 0)
        self.assertGreater(len(a.get("texts", [])), 0)

    @unittest.skipUnless(importlib.util.find_spec("PIL"), "Pillow not installed")
    def test_png_renderer_outputs_file(self):
        layout = self._sample_layout()
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "view.png"
            render_layout_to_png(layout, out)
            self.assertTrue(out.exists())
            self.assertGreater(out.stat().st_size, 0)

    @unittest.skipUnless(importlib.util.find_spec("pptx"), "python-pptx not installed")
    def test_pptx_renderer_outputs_slide_with_shapes(self):
        from pptx import Presentation

        layout = self._sample_layout()
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "view.pptx"
            export_layout_to_pptx(layout, out, "Test Export")
            self.assertTrue(out.exists())
            prs = Presentation(str(out))
            self.assertEqual(len(prs.slides), 1)
            self.assertGreater(len(prs.slides[0].shapes), 5)


if __name__ == "__main__":
    unittest.main()
